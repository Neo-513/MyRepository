import csv
import openpyxl
import re
import xlrd
import xlwt


def read(path, encoding="gbk", trace=False):
	"""读取文件
	:param path: 文件路径
	:param encoding: 文件编码
	:param trace: 是否打印信息
	:return: 文件内容集合
	"""
	filetype = path.split(".")[-1]  # 文件类型
	if trace:
		print(f"[READ]    {path}")  # 打印信息

	if filetype == "csv":
		return _read_csv(path, encoding)
	elif filetype == "xlsx":
		return _read_xlsx(path)
	elif filetype == "xls":
		return _read_xls(path)
	else:
		with open(path, mode="r", encoding="utf-8") as file:
			return file.read()


def write(path, datas, encoding="gbk", trace=False):
	"""写入文件
	:param path: 文件路径
	:param datas: 数据集合
	:param encoding: 文件编码
	:param trace: 是否打印信息
	:return:
	"""
	filetype = path.split(".")[-1]  # 文件类型
	if trace:
		print(f"[WRITE]    {path}")  # 打印信息

	if filetype == "csv":
		_write_csv(path, datas, encoding=encoding)
	elif filetype == "xlsx":
		_write_xlsx(path, datas)
	elif filetype == "xls":
		_write_xls(path, datas)
	else:
		with open(path, mode="w", encoding="utf-8") as file:
			file.write(datas)


def _read_csv(path, encoding):  # 读csv
	with open(path, mode="r", encoding=encoding) as file:
		datas = [[r for r in reader] for reader in csv.reader(file)]
		if len(datas) == 1:  # 一维列表
			datas = datas[0]
	return datas


def _write_csv(path, datas, encoding):  # 写csv
	with open(path, mode="w", encoding=encoding, newline="") as file:
		if not isinstance(datas[0], list):  # 一维列表
			datas = [datas]
		csv.writer(file).writerows(datas)


def _read_xlsx(path):  # 读xlsx
	file = openpyxl.load_workbook(path)  # 文件对象
	datas = {}  # 数据集合字典

	for sheetname in file.sheetnames:  # 遍历表名
		data = file[sheetname]  # 单张表数据集合
		datas[sheetname] = [[cell.value if cell.value else "" for cell in row] for row in data.rows]

	file.close()  # 关闭文件对象
	return datas


def _write_xlsx(path, datas):  # 写xlsx
	file = openpyxl.Workbook()  # 文件对象
	file.remove(file.active)  # 删除默认表

	for sheetname in datas:  # 遍历表名
		sheet = file.create_sheet(sheetname)  # 创建表
		data = datas[sheetname]  # 单张表数据集合
		row = len(data)  # 行数
		col = 0 if row == 0 else len(data[0])  # 列数
		for r in range(row):  # 遍历行
			for c in range(col):  # 遍历单元格
				sheet.cell(r + 1, c + 1).value = data[r][c]  # 写入单元格数据（起始索引为1）

	file.save(path)  # 保存文件
	file.close()  # 关闭文件对象


def _read_xls(path):  # 读xls
	file = xlrd.open_workbook(path)  # 文件对象
	datas = {}  # 数据集合字典

	for sheetname in file.sheet_names():  # 遍历表名
		data = file.sheet_by_name(sheetname)  # 单张表数据集合
		datas[sheetname] = [[
			int(cell.value) if re.findall("^\d+\.0*$", str(cell.value)) else cell.value for cell in row
		] for row in data.get_rows()]

	return datas


def _write_xls(path, datas):  # 写xls
	file = xlwt.Workbook()  # 文件对象

	for sheetname in datas:  # 遍历表名
		sheet = file.add_sheet(sheetname)  # 创建表
		data = datas[sheetname]  # 单张表数据集合
		row = len(data)  # 行数
		col = 0 if row == 0 else len(data[0])  # 列数
		for r in range(row):  # 遍历行
			for c in range(col):  # 遍历单元格
				sheet.write(r, c, data[r][c])  # 写入单元格数据（起始索引为0）

	file.save(path)  # 保存文件
